"""xlsx: typed cells, a named Table, a frozen header, a filter (ADR-226, ADR-274).

The oracle is openpyxl's own reader — what we write is what a reader sees — and,
at review time, Excel itself through the measurement harness.
"""

import io

import openpyxl
import pytest

from src.domains.document_generation.context import RenderContext
from src.domains.document_generation.renderers import render_document
from src.domains.document_generation.schemas import (
    DocumentType,
    SectionBlock,
    SectionedContent,
    TableSheet,
    TabularContent,
)

pytestmark = [pytest.mark.unit]
_CTX = RenderContext(language="en")


def _sheet(**overrides: object) -> TableSheet:
    base: dict[str, object] = {
        "name": "Data",
        "headers": ["City", "Date", "Amount", "Share", "Zip"],
        "rows": [
            ["Strasbourg", "2026-09-01", "1234.5", "12%", "67000"],
            ["Colmar", "2026-09-02", "-98.25", "3.5%", "01000"],
        ],
    }
    return TableSheet(**{**base, **overrides})  # type: ignore[arg-type]


def _book(*sheets: TableSheet, title: str = "Data") -> openpyxl.Workbook:
    content = TabularContent(filename_stem="d", title=title, sheets=list(sheets))
    data = render_document(DocumentType.XLSX, content, _CTX)
    return openpyxl.load_workbook(io.BytesIO(data))


class TestTheContractOfAdr226Holds:
    def test_round_trip_and_formula_neutralized(self) -> None:
        content = TabularContent(
            filename_stem="data",
            title="Data",
            sheets=[
                TableSheet(name="Feuille 1", headers=["a", "b"], rows=[["1", "=2+2"]]),
                TableSheet(name="Feuille 2", headers=["c"], rows=[["x"]]),
            ],
        )
        workbook = openpyxl.load_workbook(
            io.BytesIO(render_document(DocumentType.XLSX, content, _CTX))
        )
        assert workbook.sheetnames == ["Feuille 1", "Feuille 2"]
        sheet = workbook["Feuille 1"]
        assert sheet["A1"].value == "a"
        assert sheet["B2"].value == "'=2+2"
        assert sheet["B2"].data_type != "f"  # the probe-proven injection stays closed

    def test_sheet_titles_sanitized_and_deduplicated(self) -> None:
        content = TabularContent(
            filename_stem="data",
            title="Data",
            sheets=[
                TableSheet(name="Q1/Q2 [draft]", headers=["a"], rows=[["1"]]),
                TableSheet(name="Q1/Q2 [draft]", headers=["b"], rows=[["2"]]),
                TableSheet(name="", headers=["c"], rows=[["3"]]),
            ],
        )
        workbook = openpyxl.load_workbook(
            io.BytesIO(render_document(DocumentType.XLSX, content, _CTX))
        )
        assert len(workbook.sheetnames) == len(set(workbook.sheetnames)) == 3
        for title in workbook.sheetnames:
            assert not set(title) & set("[]:*?/\\")

    def test_requires_tabular_content(self) -> None:
        content = SectionedContent(
            filename_stem="x", title="T", blocks=[SectionBlock(kind="paragraph", text="p")]
        )
        with pytest.raises(ValueError, match="TabularContent"):
            render_document(DocumentType.XLSX, content, _CTX)


class TestColumnsAreTyped:
    def test_dates_amounts_and_percentages_become_values(self) -> None:
        sheet = _book(_sheet()).active
        assert sheet["B2"].value.date().isoformat() == "2026-09-01"
        assert sheet["B2"].number_format == "yyyy-mm-dd"
        assert sheet["C2"].value == 1234.5 and sheet["C2"].number_format == "0.00"
        assert sheet["D3"].value == pytest.approx(0.035) and sheet["D3"].number_format == "0.0%"

    def test_a_leading_zero_keeps_the_column_text(self) -> None:
        """A postal code is not a number: typing it would eat the zero."""
        sheet = _book(_sheet()).active
        assert sheet["E3"].value == "01000"

    def test_negative_numbers_are_numbers_not_defaced_strings(self) -> None:
        sheet = _book(_sheet(headers=["delta"], rows=[["-5.2"], ["3.1"]])).active
        assert sheet["A2"].value == -5.2

    def test_numeric_columns_are_right_aligned(self) -> None:
        sheet = _book(_sheet()).active
        assert sheet["C2"].alignment.horizontal == "right"
        assert sheet["A2"].alignment.horizontal in (None, "general", "left")


class TestTheWorkbookIsUsable:
    def test_table_filter_freeze_and_generated_names(self) -> None:
        workbook = _book(_sheet(), _sheet(name="Other"))
        first, second = workbook.worksheets
        assert first.freeze_panes == "A2"
        assert list(first.tables) == ["Table1"]
        assert list(second.tables) == ["Table2"]
        assert first.tables["Table1"].autoFilter is not None
        assert first.tables["Table1"].tableStyleInfo.showRowStripes

    def test_the_header_row_is_bold_and_wrapped(self) -> None:
        sheet = _book(_sheet()).active
        assert sheet["A1"].font.bold
        assert sheet["A1"].alignment.wrap_text

    def test_column_widths_stay_inside_their_bounds(self) -> None:
        from src.domains.document_generation import typography

        sheet = _book(_sheet(headers=["h"], rows=[["x" * 500]])).active
        width = sheet.column_dimensions["A"].width
        assert typography.XLSX_COLUMN_WIDTH_MIN <= width <= typography.XLSX_COLUMN_WIDTH_MAX

    def test_a_sheet_without_rows_has_no_table_but_keeps_its_header(self) -> None:
        """An empty Table reference is what Excel refuses to open."""
        sheet = _book(_sheet(rows=[])).active
        assert list(sheet.tables) == []
        assert sheet["A1"].value == "City"

    def test_duplicate_headers_are_made_unique_before_the_table(self) -> None:
        sheet = _book(_sheet(headers=["a", "a"], rows=[["1", "2"]])).active
        assert [sheet["A1"].value, sheet["B1"].value] == ["a", "a (2)"]
        assert list(sheet.tables) == ["Table1"]

    def test_the_workbook_carries_its_title(self) -> None:
        assert _book(_sheet(), title="Quarterly data").properties.title == "Quarterly data"


@pytest.mark.unit
class TestColumnWidthFollowsTheGlyphs:
    """Excel's width unit is a Latin character: an ideograph takes about two of
    them, so a CJK column sized by ``len`` shows half its content."""

    ZH = "一丁丂七丄丅丆万丈三"

    def test_a_chinese_column_is_wider_than_a_latin_one_of_equal_length(self) -> None:
        from src.domains.document_generation.renderers.xlsx import _column_width

        rows = [[self.ZH, "abcdefghij"]]
        assert _column_width("h", rows, 0) > _column_width("h", rows, 1)

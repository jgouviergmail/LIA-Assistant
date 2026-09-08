"""docx: named styles, fields, numbering, TOC, lists, tables (ADR-274).

The oracle is python-docx's own reader — plus, at review time, Word itself
through the measurement harness, which is the only thing that can say whether
the fields actually compute.
"""

import io

import docx
import pytest
from docx.oxml.ns import qn

from src.domains.document_generation.context import RenderContext
from src.domains.document_generation.renderers import render_document
from src.domains.document_generation.renderers.docx_styles import CALLOUT_STYLE
from src.domains.document_generation.schemas import (
    DocumentType,
    SectionBlock,
    SectionedContent,
    TableSheet,
)

pytestmark = [pytest.mark.unit]


def _long_document(subtitle: str = "For the board") -> SectionedContent:
    blocks: list[SectionBlock] = []
    for part in range(1, 4):
        blocks.append(SectionBlock(kind="heading", level=1, text=f"{part}. Part {part}"))
        blocks.append(SectionBlock(kind="paragraph", text="Body **strong** text."))
        blocks.append(SectionBlock(kind="heading", level=2, text="Detail"))
        blocks.append(SectionBlock(kind="numbered", items=["one", "two"]))
    blocks.append(SectionBlock(kind="quote", text="verbatim words"))
    blocks.append(SectionBlock(kind="callout", text="watch out"))
    blocks.append(
        SectionBlock(
            kind="table",
            caption="Figures",
            table=TableSheet(name="t", headers=["City", "Pop"], rows=[["A", "1"], ["B", "22"]]),
        )
    )
    return SectionedContent(filename_stem="r", title="Report", subtitle=subtitle, blocks=blocks)


def _render(content: SectionedContent, **context_kwargs: object) -> docx.document.Document:
    context = RenderContext(language="en", **context_kwargs)  # type: ignore[arg-type]
    return docx.Document(io.BytesIO(render_document(DocumentType.DOCX, content, context)))


def _xml(document: docx.document.Document) -> str:
    return document.element.xml


@pytest.fixture(autouse=True)
def _threshold(monkeypatch: pytest.MonkeyPatch) -> None:
    """Pin the long-document threshold: the test must not depend on a tunable."""
    from src.core.config import settings

    monkeypatch.setattr(settings, "document_generation_toc_min_headings", 5)


class TestTheDocumentIsBuiltFromNamedStyles:
    def test_every_block_kind_uses_its_style(self) -> None:
        document = _render(_long_document())
        styles = {paragraph.style.name for paragraph in document.paragraphs}
        assert {
            "Title",
            "Subtitle",
            "Heading 1",
            "Heading 2",
            "List Number",
            "Quote",
            CALLOUT_STYLE,
            "Caption",
        } <= styles

    def test_the_title_block_opens_the_document(self) -> None:
        document = _render(_long_document())
        assert document.paragraphs[0].text == "Report"
        assert document.paragraphs[1].text == "For the board"
        assert document.core_properties.title == "Report"
        assert document.core_properties.subject == "For the board"

    def test_the_cjk_face_is_declared_so_chinese_never_falls_back(self) -> None:
        document = _render(_long_document())
        assert 'w:eastAsia="Microsoft YaHei"' in document.styles["Normal"].element.xml

    def test_inline_emphasis_becomes_runs(self) -> None:
        document = _render(_long_document())
        strong = [
            run
            for paragraph in document.paragraphs
            for run in paragraph.runs
            if run.text == "strong"
        ]
        assert strong and strong[0].bold


class TestTheLongDocumentApparatus:
    def test_toc_numbering_and_part_breaks_arrive_together(self) -> None:
        document = _render(_long_document())
        assert 'TOC \\o "1-3" \\h \\z \\u \\n' in _xml(document)
        # "TOC Heading" also starts with "TOC ": match the entry styles only.
        entries = [
            p.text for p in document.paragraphs if p.style.name in {"TOC 1", "TOC 2", "TOC 3"}
        ]
        assert entries == [
            "1\tPart 1",
            "1.1\tDetail",
            "2\tPart 2",
            "2.1\tDetail",
            "3\tPart 3",
            "3.1\tDetail",
        ]
        # The model's own "1." prefix was stripped: Word numbers the headings.
        assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"] == [
            "Part 1",
            "Part 2",
            "Part 3",
        ]
        assert document.styles["Heading 1"].element.pPr.find(qn("w:numPr")) is not None
        breaks = [
            p
            for p in document.paragraphs
            if p.style.name == "Heading 1" and p.paragraph_format.page_break_before
        ]
        assert len(breaks) == 2  # every part but the first

    def test_the_contents_heading_is_not_counted_as_a_chapter(self) -> None:
        """TOC Heading is based on Heading 1: detached, or it becomes chapter 1."""
        document = _render(_long_document())
        properties = document.styles["TOC Heading"].element.pPr
        assert properties.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val")) == "0"
        assert properties.find(qn("w:outlineLvl")).get(qn("w:val")) == "9"

    def test_a_short_document_gets_none_of_it(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from src.core.config import settings

        monkeypatch.setattr(settings, "document_generation_toc_min_headings", 50)
        document = _render(_long_document())
        assert "TOC \\o" not in _xml(document)
        assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"] == [
            "1. Part 1",
            "2. Part 2",
            "3. Part 3",
        ]

    def test_a_plain_structure_gets_none_of_it_either(self) -> None:
        """Meeting minutes shape themselves (owner decision 2026-09-08)."""
        assert "TOC \\o" not in _xml(_render(_long_document(), structure="plain"))


class TestPagesFieldsAndTables:
    def test_the_running_head_and_the_page_footer(self) -> None:
        document = _render(_long_document())
        section = document.sections[0]
        assert section.header.paragraphs[0].text == "Report"
        footer_xml = section.footer._element.xml
        assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
        assert "Page " in section.footer.paragraphs[0].text

    def test_the_page_size_follows_the_context(self) -> None:
        assert (
            round(_render(_long_document(), page_size="letter").sections[0].page_width.inches, 1)
            == 8.5
        )
        assert round(_render(_long_document(), page_size="a4").sections[0].page_width.mm) == 210

    def test_numbered_lists_restart_at_one(self) -> None:
        """Three lists, three numbering instances: 1. 2. then 1. 2. again.

        The definitions live in ``numbering.xml``, not in the document part.
        """
        document = _render(_long_document())
        numbering_xml = document.part.numbering_part.element.xml
        assert numbering_xml.count("<w:startOverride") == 3

    def test_the_table_is_captioned_repeats_its_header_and_aligns_numbers(self) -> None:
        document = _render(_long_document())
        assert [p.text for p in document.paragraphs if p.style.name == "Caption"] == [
            "Table 1 — Figures"
        ]
        table = document.tables[0]
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        assert table.cell(1, 1).paragraphs[0].alignment == WD_ALIGN_RIGHT
        assert table.cell(1, 0).paragraphs[0].alignment in (None, 0)

    def test_the_date_line_appears_when_the_caller_gives_one(self) -> None:
        from datetime import UTC, datetime

        document = _render(
            _long_document(),
            generated_at=datetime(2026, 9, 8, 12, tzinfo=UTC),
            timezone="Europe/Berlin",
        )
        assert any("2026" in paragraph.text for paragraph in document.paragraphs[:4])


WD_ALIGN_RIGHT = 2  # WD_ALIGN_PARAGRAPH.RIGHT, as python-docx reads it back

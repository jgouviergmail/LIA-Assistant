"""The overflow oracle, shared by the pptx tests and the corpus tests (ADR-274).

python-pptx cannot tell whether text fits; the renderer decides with ``fit``,
so the oracle re-checks with the SAME formula. It is not a tautology: the
renderer chooses a size and a split, the oracle verifies the RESULT — a bug in
the splitting, in the geometry or in a placeholder lookup shows up here.
"""

from __future__ import annotations

from typing import Any

from pptx.enum.shapes import PP_PLACEHOLDER

from src.domains.document_generation.fit import BULLET_INDENT_PT, TextFrame, fits
from src.domains.document_generation.renderers.pptx_geometry import EMU_PER_PT

_TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
#: Size assumed for a run that carries none (inherited from the master).
_DEFAULT_SIZE_PT = 18.0


def assert_nothing_overflows(presentation: Any) -> None:
    """Every text frame of every slide stays inside the estimator's budget.

    Args:
        presentation: A python-pptx presentation, freshly read back.

    Raises:
        AssertionError: Naming the shape and its first paragraphs.
    """
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            placeholder_type = shape.placeholder_format.type if shape.is_placeholder else None
            if placeholder_type == PP_PLACEHOLDER.SLIDE_NUMBER:
                continue
            indent = 0.0 if placeholder_type in _TITLE_TYPES else BULLET_INDENT_PT
            frame = TextFrame(shape.width / EMU_PER_PT, shape.height / EMU_PER_PT, indent_pt=indent)
            paragraphs = [paragraph.text for paragraph in shape.text_frame.paragraphs]
            sizes = [
                run.font.size.pt
                for paragraph in shape.text_frame.paragraphs
                for run in paragraph.runs
                if run.font.size is not None
            ]
            size = max(sizes) if sizes else _DEFAULT_SIZE_PT
            assert fits(
                paragraphs, size, frame
            ), f"{shape.name} overflows at {size}pt: {paragraphs[:2]}"

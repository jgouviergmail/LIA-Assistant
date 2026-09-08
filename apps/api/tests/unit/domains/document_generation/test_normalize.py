"""Every incoherence is repaired, none is refused (ADR-184, ADR-274).

A bound in the schema fails the whole document for a detail (ADR-269); a
renderer that trusted the model's shape would draw a table with no rows or a
comparison with one side. Everything a writer can plausibly produce is put in
canonical form here, mechanically, and nothing is reported as a defect.
"""

import io

import pytest

from src.domains.document_generation.context import RenderContext
from src.domains.document_generation.normalize import (
    HeadingNumberer,
    document_is_numbered,
    normalize_content,
    normalize_sectioned,
    normalize_slides,
    normalize_tabular,
)
from src.domains.document_generation.schemas import (
    SectionBlock,
    SectionedContent,
    Slide,
    SlideColumn,
    SlideContent,
    TableSheet,
    TabularContent,
)

pytestmark = [pytest.mark.unit]
_CTX = RenderContext(language="en")


def _doc(*blocks: SectionBlock) -> SectionedContent:
    return SectionedContent(filename_stem="d", title="T", blocks=list(blocks))


def _blocks(*blocks: SectionBlock, context: RenderContext = _CTX) -> list[SectionBlock]:
    return normalize_sectioned(_doc(*blocks), context).blocks


class TestSectionedRepairs:
    def test_levels_are_clamped_and_empty_blocks_dropped(self) -> None:
        out = _blocks(
            SectionBlock(kind="heading", level=9, text="H"),
            SectionBlock(kind="heading", level=0, text="  "),
            SectionBlock(kind="paragraph", text=" "),
            SectionBlock(kind="bullets", items=["", "  "]),
            SectionBlock(kind="table"),
        )
        assert [(b.kind, b.level, b.text) for b in out] == [("heading", 4, "H")]

    def test_a_level_zero_heading_keeps_its_text(self) -> None:
        assert _blocks(SectionBlock(kind="heading", level=0, text="H"))[0].level == 1

    def test_markdown_that_leaked_into_a_paragraph_becomes_its_block(self) -> None:
        out = _blocks(
            SectionBlock(kind="paragraph", text="- a\n- b\n* c"),
            SectionBlock(kind="paragraph", text="1. one\n2) two"),
            SectionBlock(kind="paragraph", text="### Inner"),
            SectionBlock(kind="bullets", text="just text", items=[]),
        )
        assert [b.kind for b in out] == ["bullets", "numbered", "heading", "paragraph"]
        assert out[0].items == ["a", "b", "c"]
        assert out[1].items == ["one", "two"]
        assert out[2].level == 3 and out[2].text == "Inner"
        assert out[3].text == "just text"

    def test_a_paragraph_that_merely_mentions_a_dash_is_left_alone(self) -> None:
        """The falsifier: only a FULL list becomes a list."""
        text = "Nous avons vu - comme prévu - trois options."
        assert _blocks(SectionBlock(kind="paragraph", text=text))[0].kind == "paragraph"

    def test_captions_only_survive_on_tables(self) -> None:
        out = _blocks(
            SectionBlock(
                kind="table",
                caption="  Key figures  ",
                table=TableSheet(name="t", headers=["a", "a"], rows=[["1"]]),
            ),
            SectionBlock(kind="paragraph", text="p", caption="ignored"),
        )
        assert out[0].table is not None and out[0].table.headers == ["a", "a (2)"]
        assert out[0].caption == "Key figures"
        assert out[1].caption == ""

    def test_whitespace_is_collapsed_everywhere(self) -> None:
        out = _blocks(SectionBlock(kind="paragraph", text="a\n\n   b\tc"))
        assert out[0].text == "a b c"

    def test_a_document_of_nothing_still_renders_its_title(self) -> None:
        """min_length=1 keeps the schema honest; the renderer still needs a block."""
        out = normalize_sectioned(_doc(SectionBlock(kind="paragraph", text="  ")), _CTX)
        assert len(out.blocks) == 1 and out.blocks[0].text == "T"


class TestTheLongDocumentApparatus:
    def test_heading_numbers_are_stripped_only_when_the_renderer_numbers(self, monkeypatch) -> None:
        from src.core.config import settings

        monkeypatch.setattr(settings, "document_generation_toc_min_headings", 2)
        blocks = (
            SectionBlock(kind="heading", level=1, text="1. Context"),
            SectionBlock(kind="heading", level=2, text="1.1) Market"),
        )
        numbered = normalize_sectioned(_doc(*blocks), _CTX)
        assert [b.text for b in numbered.blocks] == ["Context", "Market"]
        assert document_is_numbered(numbered, _CTX)

        plain_ctx = RenderContext(language="en", structure="plain")
        plain = normalize_sectioned(_doc(*blocks), plain_ctx)
        assert [b.text for b in plain.blocks] == ["1. Context", "1.1) Market"]
        assert not document_is_numbered(plain, plain_ctx)

    def test_a_short_document_is_not_numbered(self, monkeypatch) -> None:
        from src.core.config import settings

        monkeypatch.setattr(settings, "document_generation_toc_min_headings", 5)
        doc = normalize_sectioned(_doc(SectionBlock(kind="heading", text="Only one")), _CTX)
        assert not document_is_numbered(doc, _CTX)

    def test_a_heading_that_is_only_a_number_keeps_its_text(self) -> None:
        """Stripping must never empty a heading."""
        from src.core.config import settings

        with pytest.MonkeyPatch.context() as patch:
            patch.setattr(settings, "document_generation_toc_min_headings", 2)
            out = _blocks(
                SectionBlock(kind="heading", level=1, text="1."),
                SectionBlock(kind="heading", level=1, text="2. Real"),
            )
        assert out[0].text == "1."
        assert out[1].text == "Real"

    def test_the_numberer_counts_and_resets(self) -> None:
        numberer = HeadingNumberer()
        assert [numberer.number(level) for level in (1, 2, 2, 3, 1, 4)] == [
            "1",
            "1.1",
            "1.2",
            "1.2.1",
            "2",
            "",
        ]


def _slides(*slides: Slide) -> list[Slide]:
    content = SlideContent(filename_stem="d", title="D", slides=list(slides))
    return normalize_slides(content, "en").slides


class TestEffectiveSlideKinds:
    def test_the_payload_decides_the_kind(self) -> None:
        table = TableSheet(name="t", headers=["a"], rows=[["1"]])
        assert [s.kind for s in _slides(Slide(title="t", bullets=["x"]))] == ["content"]
        assert [s.kind for s in _slides(Slide(title="t", table=table))] == ["table"]
        assert [s.kind for s in _slides(Slide(title="t", bullets=["x"], table=table))] == [
            "content",
            "table",
        ]
        assert [s.kind for s in _slides(Slide(title="t", kind="table"))] == ["content"]
        assert [s.kind for s in _slides(Slide(title="t"))] == ["content"]

    def test_a_section_with_bullets_becomes_a_divider_then_a_content_slide(self) -> None:
        out = _slides(Slide(title="Part", kind="section", subtitle="tag", bullets=["a"], notes="n"))
        assert [(s.kind, s.subtitle, s.bullets) for s in out] == [
            ("section", "tag", []),
            ("content", "tag", ["a"]),
        ]
        assert out[0].notes == "n" and out[1].notes == ""

    def test_comparison_with_one_two_or_three_columns(self) -> None:
        two = [
            SlideColumn(heading="A", bullets=["a1"]),
            SlideColumn(heading="B", bullets=["b1", "b2"]),
        ]
        assert [s.kind for s in _slides(Slide(title="t", kind="comparison", columns=two))] == [
            "comparison"
        ]
        # Declared content, but the payload says comparison.
        assert [s.kind for s in _slides(Slide(title="t", columns=two))] == ["comparison"]

        one = _slides(Slide(title="t", kind="comparison", columns=two[:1]))
        assert one[0].kind == "content" and one[0].bullets == ["a1"]

        three = _slides(
            Slide(
                title="t",
                kind="comparison",
                columns=[*two, SlideColumn(heading="C", bullets=["c1"])],
            )
        )
        assert three[0].kind == "table"
        assert three[0].table is not None
        assert three[0].table.headers == ["A", "B", "C"]
        assert three[0].table.rows == [["a1", "b1", "c1"], ["", "b2", ""]]

    def test_empty_text_is_dropped_and_notes_stay_on_the_first_slide(self) -> None:
        out = _slides(Slide(title="  t  ", bullets=["", " ", "real"], notes=" n "))
        assert out[0].title == "t" and out[0].bullets == ["real"] and out[0].notes == "n"

    def test_a_slide_carrying_an_empty_table_falls_back_to_content(self) -> None:
        empty = TableSheet(name="t", headers=[], rows=[])
        assert [s.kind for s in _slides(Slide(title="t", kind="table", table=empty))] == ["content"]


class TestTabularAndDispatch:
    def test_sheets_are_normalized_and_a_workbook_always_has_one(self) -> None:
        out = normalize_tabular(
            TabularContent(
                filename_stem="t",
                title="Title",
                sheets=[TableSheet(name="s", headers=[], rows=[])],
            ),
            "en",
        )
        assert len(out.sheets) == 1 and out.sheets[0].headers == ["Title"]

    def test_normalize_content_dispatches_on_the_family(self) -> None:
        sectioned = normalize_content(_doc(SectionBlock(kind="paragraph", text="p")), _CTX)
        assert isinstance(sectioned, SectionedContent)
        deck = normalize_content(
            SlideContent(filename_stem="d", title="D", slides=[Slide(title="t")]), _CTX
        )
        assert isinstance(deck, SlideContent)
        book = normalize_content(
            TabularContent(
                filename_stem="t",
                title="T",
                sheets=[TableSheet(name="s", headers=["a"], rows=[["1"]])],
            ),
            _CTX,
        )
        assert isinstance(book, TabularContent)

    def test_the_dispatch_covers_every_family(self) -> None:
        """ADR-085: a family the dispatch cannot canonicalise would reach a renderer raw."""
        from src.domains.document_generation.normalize import _NORMALIZERS
        from src.domains.document_generation.schemas import SCHEMA_BY_DOC_TYPE

        assert set(_NORMALIZERS) == set(SCHEMA_BY_DOC_TYPE.values())

    def test_normalizing_twice_changes_nothing(self) -> None:
        """Canonical means fixed: a renderer may normalize defensively."""
        once = normalize_content(
            _doc(
                SectionBlock(kind="paragraph", text="- a\n- b"),
                SectionBlock(kind="heading", level=9, text="H"),
            ),
            _CTX,
        )
        twice = normalize_content(once, _CTX)
        assert twice == once


@pytest.mark.unit
class TestNothingIsEmptyEnoughToBreakARenderer:
    """A model can answer with whitespace; a renderer must still produce a file."""

    @pytest.mark.parametrize(
        ("doc_type", "content"),
        [
            (
                "docx",
                SectionedContent(
                    filename_stem="x",
                    title="   ",
                    blocks=[SectionBlock(kind="paragraph", text="  ")],
                ),
            ),
            (
                "pdf",
                SectionedContent(
                    filename_stem="x", title="", blocks=[SectionBlock(kind="paragraph", text="")]
                ),
            ),
            ("pptx", SlideContent(filename_stem="x", title="  ", slides=[Slide(title="   ")])),
            (
                "xlsx",
                TabularContent(
                    filename_stem="x",
                    title="",
                    sheets=[TableSheet(name="", headers=[], rows=[])],
                ),
            ),
            (
                "md",
                SectionedContent(
                    filename_stem="x", title=" ", blocks=[SectionBlock(kind="bullets", items=[" "])]
                ),
            ),
        ],
    )
    def test_a_document_of_whitespace_still_renders(self, doc_type: str, content: object) -> None:
        from src.domains.document_generation.renderers import render_document
        from src.domains.document_generation.schemas import DocumentType

        data = render_document(DocumentType(doc_type), content, _CTX)  # type: ignore[arg-type]
        assert isinstance(data, bytes)


@pytest.mark.unit
class TestAListTheModelPutInTheWrongField:
    """``items`` is optional, so a writer can put its list in ``text`` instead."""

    def test_markdown_in_the_text_of_a_list_block_is_recovered(self) -> None:
        """Falling back to a paragraph merged two bullets into one line of prose
        with a stray dash — and made normalization non-idempotent, because the
        second pass DID recognise the markdown the first pass had flattened."""
        block = SectionBlock(kind="bullets", items=[], text="- alpha\n- beta")
        canonical = normalize_sectioned(_doc(block), _CTX)
        assert [(b.kind, b.items) for b in canonical.blocks] == [("bullets", ["alpha", "beta"])]

    def test_an_ordered_list_in_the_text_field_keeps_its_kind(self) -> None:
        block = SectionBlock(kind="numbered", items=[], text="1. un\n2. deux")
        canonical = normalize_sectioned(_doc(block), _CTX)
        assert [(b.kind, b.items) for b in canonical.blocks] == [("numbered", ["un", "deux"])]

    def test_plain_text_in_a_list_block_still_becomes_a_paragraph(self) -> None:
        block = SectionBlock(kind="bullets", items=[], text="Une phrase ordinaire.")
        canonical = normalize_sectioned(_doc(block), _CTX)
        assert [(b.kind, b.text) for b in canonical.blocks] == [
            ("paragraph", "Une phrase ordinaire.")
        ]

    def test_the_repair_is_idempotent(self) -> None:
        block = SectionBlock(kind="bullets", items=[], text="- alpha\n- beta")
        once = normalize_sectioned(_doc(block), _CTX)
        assert normalize_sectioned(once, _CTX) == once


@pytest.mark.unit
class TestTheFallbackSheetIsSanitizedLikeEveryOther:
    def test_a_workbook_with_nothing_gets_a_named_column(self) -> None:
        """The fallback sheet used to bypass ``sanitize_headers``, so a workbook
        with no title carried a column with no name."""
        from src.domains.document_generation.normalize import normalize_tabular
        from src.domains.document_generation.schemas import TableSheet, TabularContent

        empty = TabularContent(
            filename_stem="x", title="", sheets=[TableSheet(name="", headers=[], rows=[])]
        )
        assert normalize_tabular(empty, "en").sheets[0].headers == ["Column 1"]

    def test_a_titled_workbook_keeps_its_title_as_the_column(self) -> None:
        from src.domains.document_generation.normalize import normalize_tabular
        from src.domains.document_generation.schemas import TableSheet, TabularContent

        empty = TabularContent(
            filename_stem="x", title="Données", sheets=[TableSheet(name="", headers=[], rows=[])]
        )
        assert normalize_tabular(empty, "fr").sheets[0].headers == ["Données"]


def _text_of(doc_type: str, data: bytes) -> str:
    """Everything a reader of that format can see, as one string.

    The oracle a renderer cannot fake: python-pptx accepted a NUL byte where
    python-docx refused it, so "no exception" proves nothing on its own.
    """
    if doc_type in ("md", "txt", "csv"):
        return data.decode("utf-8-sig")
    if doc_type == "docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if doc_type == "pptx":
        import pptx

        presentation = pptx.Presentation(io.BytesIO(data))
        return "\n".join(
            shape.text_frame.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
    if doc_type == "xlsx":
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(data))
        return "\n".join(
            str(cell.value)
            for sheet in book.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
    import fitz

    with fitz.open(stream=data, filetype="pdf") as document:
        return "\n".join(page.get_text() for page in document)


@pytest.mark.unit
class TestAControlCharacterNeverReachesARenderer:
    """One repair in the canonical form covers all seven formats."""

    BAD = "avant\x00milieu\x07fin"

    @pytest.mark.parametrize("doc_type", ["docx", "pdf", "pptx", "xlsx", "md", "txt", "csv"])
    def test_every_format_renders_text_carrying_control_characters(self, doc_type: str) -> None:
        from src.domains.document_generation.renderers import render_document
        from src.domains.document_generation.schemas import (
            DocumentType,
            SectionBlock,
            SectionedContent,
            Slide,
            SlideContent,
            TableSheet,
            TabularContent,
        )

        family = {
            "docx": "sectioned",
            "pdf": "sectioned",
            "md": "sectioned",
            "txt": "sectioned",
            "pptx": "slides",
            "xlsx": "tabular",
            "csv": "tabular",
        }[doc_type]
        if family == "sectioned":
            content = SectionedContent(
                filename_stem="x",
                title=self.BAD,
                blocks=[SectionBlock(kind="paragraph", text=self.BAD)],
            )
        elif family == "slides":
            content = SlideContent(
                filename_stem="x",
                title=self.BAD,
                slides=[Slide(title=self.BAD, bullets=[self.BAD], notes=self.BAD)],
            )
        else:
            content = TabularContent(
                filename_stem="x",
                title=self.BAD,
                sheets=[TableSheet(name=self.BAD, headers=[self.BAD], rows=[[self.BAD]])],
            )
        data = render_document(DocumentType(doc_type), content, _CTX)
        assert data
        assert "\x00" not in _text_of(doc_type, data)

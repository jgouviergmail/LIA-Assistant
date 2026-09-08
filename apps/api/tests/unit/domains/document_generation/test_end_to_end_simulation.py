"""The whole chain, from the tool to the bytes on disk (ADR-226, ADR-274).

Every other test exercises one seam. This one runs the real path — tool guards,
service, prompt rendering, normalization, renderer, attachment row, card — with
only the LLM and the database replaced, and then OPENS the produced file with
the format's own reader. It is the simulation a unit test usually skips, and
it is where a wiring mistake between two green modules would show.
"""

import io
import uuid
from contextlib import asynccontextmanager
from unittest.mock import AsyncMock, MagicMock, patch

import docx
import openpyxl
import pptx
import pytest

from src.domains.document_generation.schemas import (
    SectionBlock,
    SectionedContent,
    Slide,
    SlideColumn,
    SlideContent,
    TableSheet,
    TabularContent,
)

pytestmark = [pytest.mark.unit]
_ATTACHMENT_ID = uuid.UUID("00000000-0000-0000-0000-0000000000ee")


class _FakeAttachment:
    def __init__(self, expires_at: object) -> None:
        self.id = _ATTACHMENT_ID
        self.expires_at = expires_at


class _FakeRepo:
    created: dict = {}

    def __init__(self, db: object) -> None:
        self._db = db

    async def create(self, payload: dict) -> _FakeAttachment:
        type(self).created = dict(payload)
        return _FakeAttachment(expires_at=payload["expires_at"])


@asynccontextmanager
async def _fake_db():
    class _Db:
        async def commit(self) -> None:
            return None

    yield _Db()


def _user(language: str = "fr", timezone: str = "Europe/Paris") -> MagicMock:
    user = MagicMock()
    user.language = language
    user.timezone = timezone
    return user


REPORT = SectionedContent(
    filename_stem="rapport-trimestriel",
    title="Rapport trimestriel",
    subtitle="À l'attention du comité",
    blocks=[SectionBlock(kind="paragraph", text="Introduction au **trimestre**.")]
    + [
        block
        for part in range(1, 4)
        for block in (
            SectionBlock(kind="heading", level=1, text=f"{part}. Partie {part}"),
            SectionBlock(kind="paragraph", text="Analyse détaillée. " * 20),
            SectionBlock(kind="heading", level=2, text="Détail"),
            SectionBlock(kind="numbered", items=["constater", "chiffrer", "décider"]),
        )
    ]
    + [
        SectionBlock(kind="callout", text="Les chiffres de mars sont provisoires."),
        SectionBlock(
            kind="table",
            caption="Chiffres clés",
            table=TableSheet(
                name="t",
                headers=["Ville", "Montant", "Part"],
                rows=[["Strasbourg", "1234.5", "12%"], ["Colmar", "980.0", "9.5%"]],
            ),
        ),
    ],
)

DECK = SlideContent(
    filename_stem="presentation",
    title="Analyse",
    subtitle="Comité",
    slides=[
        Slide(title="1. Contexte", kind="section", subtitle="Le marché"),
        Slide(title="Les coûts montent", bullets=["Hausse de **12 %**", "Trois fournisseurs"]),
        Slide(
            title="A contre B",
            kind="comparison",
            columns=[
                SlideColumn(heading="A", bullets=["moins cher"]),
                SlideColumn(heading="B", bullets=["plus rapide"]),
            ],
        ),
    ],
)

WORKBOOK = TabularContent(
    filename_stem="donnees",
    title="Données",
    sheets=[
        TableSheet(
            name="Ventes",
            headers=["Ville", "Date", "Montant"],
            rows=[["Strasbourg", "2026-09-01", "1234.5"], ["Colmar", "2026-09-02", "980.0"]],
        )
    ],
)


async def _run_tool(doc_type: str, content: object, tmp_path, monkeypatch) -> tuple[object, bytes]:
    """Drive the real tool, return its output and the bytes it wrote."""
    from src.core.config import settings as app_settings
    from src.domains.agents.tools import document_generation_tools as tool_module
    from src.domains.document_generation import service as svc
    from tests.unit.domains.agents.tools.test_document_generation_tools import (
        SETTINGS_PATCH_PATH,
        _fake_settings,
        _runtime,
        _wrapper_settings,
    )

    monkeypatch.setattr(app_settings, "attachments_storage_path", str(tmp_path))
    _FakeRepo.created = {}
    with (
        patch.object(tool_module, "settings", _fake_settings()),
        patch(SETTINGS_PATCH_PATH, return_value=_wrapper_settings()),
        patch.object(tool_module, "_load_user", AsyncMock(return_value=_user())),
        patch.object(svc, "AttachmentRepository", _FakeRepo),
        patch.object(svc, "get_db_context", _fake_db),
        patch.object(svc, "_call_document_llm", AsyncMock(return_value=content)),
    ):
        result = await tool_module.generate_document.coroutine(
            instructions="peu importe, le modèle est simulé",
            doc_type=doc_type,
            runtime=_runtime(),
        )
    stored = tmp_path / _FakeRepo.created["file_path"]
    return result, stored.read_bytes()


@pytest.fixture(autouse=True)
def _drain_pending_cards() -> None:
    """The pending store is module-level and keyed by conversation: it holds
    cards until the SSE done chunk drains them, so each case starts empty."""
    from src.domains.document_generation.document_store import (
        get_and_clear_pending_documents,
    )

    get_and_clear_pending_documents("conv1")


class TestTheWholeChain:
    async def test_a_report_reaches_disk_as_a_crafted_docx(self, tmp_path, monkeypatch) -> None:
        result, data = await _run_tool("docx", REPORT, tmp_path, monkeypatch)
        assert result.success is True
        assert result.structured_data["filename"] == "rapport-trimestriel.docx"

        document = docx.Document(io.BytesIO(data))
        styles = {paragraph.style.name for paragraph in document.paragraphs}
        assert {"Title", "Subtitle", "Heading 1", "List Number", "Caption"} <= styles
        # The reader's own language reached the renderer through the user row.
        assert "Sommaire" in [paragraph.text for paragraph in document.paragraphs]
        assert "Tableau 1 — Chiffres clés" in [p.text for p in document.paragraphs]
        # A dated title block, in the reader's timezone.
        assert any("2026" in paragraph.text for paragraph in document.paragraphs[:4])
        assert document.sections[0].header.paragraphs[0].text == "Rapport trimestriel"

    async def test_a_deck_reaches_disk_as_a_16_9_pptx(self, tmp_path, monkeypatch) -> None:
        from .pptx_oracles import assert_nothing_overflows

        result, data = await _run_tool("pptx", DECK, tmp_path, monkeypatch)
        assert result.success is True
        presentation = pptx.Presentation(io.BytesIO(data))
        assert presentation.slide_width > presentation.slide_height
        assert [slide.slide_layout.name for slide in presentation.slides] == [
            "Title Slide",
            "Section Header",
            "Title and Content",
            "Comparison",
        ]
        assert_nothing_overflows(presentation)

    async def test_a_workbook_reaches_disk_with_typed_columns(self, tmp_path, monkeypatch) -> None:
        result, data = await _run_tool("xlsx", WORKBOOK, tmp_path, monkeypatch)
        assert result.success is True
        sheet = openpyxl.load_workbook(io.BytesIO(data)).active
        assert sheet["C2"].value == 1234.5  # a number, not a string
        assert sheet["B2"].value.date().isoformat() == "2026-09-01"
        assert sheet.freeze_panes == "A2"
        assert list(sheet.tables) == ["Table1"]

    async def test_the_card_is_queued_once_with_the_right_mime(self, tmp_path, monkeypatch) -> None:
        from src.domains.document_generation.document_store import (
            get_and_clear_pending_documents,
        )

        result, _data = await _run_tool("pdf", REPORT, tmp_path, monkeypatch)
        assert result.success is True
        assert _FakeRepo.created["mime_type"] == "application/pdf"
        assert _FakeRepo.created["content_type"] == "document"
        pending = get_and_clear_pending_documents("conv1")
        assert len(pending) == 1
        assert pending[0].filename.endswith(".pdf")

    async def test_a_chinese_reader_gets_a_chinese_apparatus(self, tmp_path, monkeypatch) -> None:
        """The labels the renderer writes follow the READER, not the deployment."""
        from src.core.config import settings as app_settings
        from src.domains.agents.tools import document_generation_tools as tool_module
        from src.domains.document_generation import service as svc
        from tests.unit.domains.agents.tools.test_document_generation_tools import (
            SETTINGS_PATCH_PATH,
            _fake_settings,
            _runtime,
            _wrapper_settings,
        )

        monkeypatch.setattr(app_settings, "attachments_storage_path", str(tmp_path))
        _FakeRepo.created = {}
        with (
            patch.object(tool_module, "settings", _fake_settings()),
            patch(SETTINGS_PATCH_PATH, return_value=_wrapper_settings()),
            patch.object(tool_module, "_load_user", AsyncMock(return_value=_user(language="zh"))),
            patch.object(svc, "AttachmentRepository", _FakeRepo),
            patch.object(svc, "get_db_context", _fake_db),
            patch.object(svc, "_call_document_llm", AsyncMock(return_value=REPORT)),
        ):
            await tool_module.generate_document.coroutine(
                instructions="x", doc_type="docx", runtime=_runtime()
            )
        data = (tmp_path / _FakeRepo.created["file_path"]).read_bytes()
        document = docx.Document(io.BytesIO(data))
        texts = [paragraph.text for paragraph in document.paragraphs]
        assert "目录" in texts  # the contents heading, in the reader's language

    async def test_a_chinese_deck_is_split_rather_than_clipped(self, tmp_path, monkeypatch) -> None:
        """The path where a full-width glyph decides the layout: PowerPoint
        measured 2 overflows on this shape, one of 124.1 pt, when an ideograph
        was counted at the Latin width (ADR-274)."""
        from src.core.config import settings as app_settings
        from src.domains.agents.tools import document_generation_tools as tool_module
        from src.domains.document_generation import service as svc
        from tests.unit.domains.agents.tools.test_document_generation_tools import (
            SETTINGS_PATCH_PATH,
            _fake_settings,
            _runtime,
            _wrapper_settings,
        )

        from .pptx_oracles import assert_nothing_overflows

        dense = (
            "这是一个关于季度业绩的详细说明我们需要在下个季度显著提高团队的整体效率并同时降低运营成本"
            * 3
        )
        deck = SlideContent(
            filename_stem="zh-deck",
            title="\u5b63\u5ea6\u62a5\u544a",
            slides=[Slide(title="\u8981\u70b9", bullets=[dense] * 9)],
        )
        monkeypatch.setattr(app_settings, "attachments_storage_path", str(tmp_path))
        _FakeRepo.created = {}
        with (
            patch.object(tool_module, "settings", _fake_settings()),
            patch(SETTINGS_PATCH_PATH, return_value=_wrapper_settings()),
            patch.object(tool_module, "_load_user", AsyncMock(return_value=_user(language="zh"))),
            patch.object(svc, "AttachmentRepository", _FakeRepo),
            patch.object(svc, "get_db_context", _fake_db),
            patch.object(svc, "_call_document_llm", AsyncMock(return_value=deck)),
        ):
            result = await tool_module.generate_document.coroutine(
                instructions="x", doc_type="pptx", runtime=_runtime()
            )
        assert result.success is True
        data = (tmp_path / _FakeRepo.created["file_path"]).read_bytes()
        presentation = pptx.Presentation(io.BytesIO(data))
        # Nine dense Chinese bullets cannot share one slide: cover + parts.
        assert len(presentation.slides) >= 3
        assert_nothing_overflows(presentation)
        written = "".join(
            shape.text_frame.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        assert written.count(dense) == 9  # nothing was clipped or dropped

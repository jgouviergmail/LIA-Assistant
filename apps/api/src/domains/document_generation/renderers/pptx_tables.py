"""Native PowerPoint tables, chunked so nothing ever runs off a slide (ADR-274).

A table is placed on a "Title Only" slide, in the template's own greyscale
style, with its header repeated on every part: a reader of part 3 must still
know what the columns are.
"""

from __future__ import annotations

from typing import Any

from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from src.domains.document_generation import typography
from src.domains.document_generation.fit import (
    TABLE_ROW_HEIGHT_PT,
    TextFrame,
    display_columns,
    rows_per_slide,
    table_font_size,
)
from src.domains.document_generation.renderers import pptx_geometry as geometry
from src.domains.document_generation.schemas import TableSheet
from src.domains.document_generation.tables import ColumnType, chunk_rows, infer_column_types

#: Shortest column width used when sizing by content, in characters.
_MIN_WEIGHT = 4


def table_area() -> TextFrame:
    """The room a table has on a slide, in points."""
    return TextFrame(
        geometry.TABLE_WIDTH_IN * 72,
        geometry.TABLE_AREA_HEIGHT_IN * 72,
        indent_pt=0,
    )


def split_rows(sheet: TableSheet) -> list[list[list[str]]]:
    """The sheet's rows in slide-sized chunks (at least one chunk)."""
    per_slide = rows_per_slide(len(sheet.headers), table_area())
    return chunk_rows(sheet.rows, per_slide) or [[]]


def _cell(cell: Any, text: str, font_pt: int, *, bold: bool, right: bool) -> None:
    cell.text = text
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.font.size = Pt(font_pt)
    paragraph.font.bold = bold
    if right:
        paragraph.alignment = PP_ALIGN.RIGHT


def _column_weights(sheet: TableSheet) -> list[float]:
    """Relative widths, in average Latin characters — never in code points.

    Ten ideographs are twice as wide as ten letters: counting them equal
    starves the CJK column, whose cells then wrap and grow the row past the
    height reserved for the table.
    """
    return [
        max(
            display_columns(header),
            *(display_columns(row[index]) for row in sheet.rows),
            _MIN_WEIGHT,
        )
        for index, header in enumerate(sheet.headers)
    ]


def add_table(slide: Any, sheet: TableSheet, types: list[ColumnType]) -> Any:
    """Place one chunk as a native table, header row included.

    Args:
        slide: The "Title Only" slide receiving the table.
        sheet: Headers plus the rows of THIS chunk.
        types: Column types of the WHOLE table, so alignment stays consistent
            across the chunks.

    Returns:
        The graphic frame holding the table.
    """
    font_pt = table_font_size(len(sheet.headers))
    row_height = Pt(TABLE_ROW_HEIGHT_PT[font_pt])
    shape = slide.shapes.add_table(
        len(sheet.rows) + 1,
        len(sheet.headers),
        Inches(geometry.TABLE_LEFT_IN),
        Inches(geometry.TABLE_TOP_IN),
        Inches(geometry.TABLE_WIDTH_IN),
        row_height * (len(sheet.rows) + 1),
    )
    table = shape.table
    style_id = shape._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId"))
    style_id.text = typography.PPTX_TABLE_STYLE_ID
    table.first_row = True
    table.horz_banding = True
    table.first_col = False

    weights = _column_weights(sheet)
    total = sum(weights)
    for index, weight in enumerate(weights):
        table.columns[index].width = int(Inches(geometry.TABLE_WIDTH_IN) * weight / total)
    for column, header in enumerate(sheet.headers):
        _cell(table.cell(0, column), header, font_pt, bold=True, right=types[column].numeric)
    for row_index, row in enumerate(sheet.rows, start=1):
        for column, value in enumerate(row):
            _cell(
                table.cell(row_index, column),
                value,
                font_pt,
                bold=False,
                right=types[column].numeric,
            )
    return shape


def column_types(sheet: TableSheet) -> list[ColumnType]:
    """The whole table's column types, computed once for every chunk."""
    return infer_column_types(sheet.headers, sheet.rows)

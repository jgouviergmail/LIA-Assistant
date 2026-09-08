"""pptx: a layout per kind, slide numbers, text that always fits (ADR-226, ADR-274).

The deck is 16:9 LANDSCAPE and uses the template's own layouts — Section
Header for a part opener, Comparison for two sides, Title Only for data — so a
reader gets the deck PowerPoint would have made, not eleven identical bullet
slides.

Nothing overflows: every body is measured by ``fit`` before it is placed,
shrunk to the floor, then SPLIT into "Title (2/3)" slides. A bullet that does
not fit alone is cut at sentence boundaries — never clipped.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Sequence
from typing import Any

from pptx.dml.color import RGBColor
from pptx.util import Pt

from src.domains.document_generation import typography
from src.domains.document_generation.context import RenderContext
from src.domains.document_generation.fit import (
    BODY_SIZE_FLOOR_PT,
    BODY_SIZE_STEP_PT,
    SlidePlan,
    fit_title,
    fits,
    plan_body,
)
from src.domains.document_generation.inline import parse_inline
from src.domains.document_generation.renderers import pptx_geometry as geometry
from src.domains.document_generation.renderers import pptx_tables
from src.domains.document_generation.schemas import (
    SLIDE_KINDS,
    DocumentContent,
    Slide,
    SlideContent,
)


def _continued(title: str, index: int, total: int) -> str:
    """« Cities (2/4) » — a numeric label, so it needs no translation."""
    return title if total == 1 else f"{title} ({index}/{total})"


def _set_title(shape: Any, text: str) -> None:
    """Write a title at the largest size that fits its frame."""
    size, _lines = fit_title(text, geometry.frame_of(shape, indent_pt=0))
    frame = shape.text_frame
    frame.text = ""
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)


def _fit_single(lines: Sequence[str], shape: Any, base_pt: int) -> int:
    """The largest size from ``base_pt`` down to the floor at which ``lines`` fit."""
    frame = geometry.frame_of(shape)
    size = base_pt
    while size > BODY_SIZE_FLOOR_PT and not fits(lines, size, frame):
        size -= BODY_SIZE_STEP_PT
    return size


def _fill(
    text_frame: Any, items: Sequence[str], size_pt: float, *, color: str | None = None
) -> None:
    """Write items as paragraphs, carrying the model's inline emphasis."""
    text_frame.text = ""
    for index, item in enumerate(items):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        for span in parse_inline(item):
            run = paragraph.add_run()
            run.text = span.text
            run.font.size = Pt(size_pt)
            run.font.bold = span.bold or None
            run.font.italic = span.italic or None
            if color is not None:
                run.font.color.rgb = RGBColor.from_string(color)


class _Deck:
    """Per-deck state: the presentation, its layouts and the two body frames."""

    def __init__(self, content: SlideContent, context: RenderContext) -> None:
        self.content = content
        self.context = context
        self.presentation = geometry.base_presentation()
        self.layouts = self.presentation.slide_layouts
        # Layout placeholders carry the geometry every slide of that layout gets,
        # so the frames are read once rather than probed slide by slide.
        self.body_frame = geometry.frame_of(
            geometry.placeholder(self.layouts[geometry.LAYOUT_CONTENT], geometry.BODY_IDX)
        )
        self.half_frame = geometry.frame_of(
            geometry.placeholder(self.layouts[geometry.LAYOUT_COMPARISON], 2)
        )

    def _new(self, layout_index: int, title: str, notes: str = "") -> Any:
        layout = self.layouts[layout_index]
        slide = self.presentation.slides.add_slide(layout)
        _set_title(slide.shapes.title, title)
        geometry.add_slide_number(slide, layout)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    # -- kinds ------------------------------------------------------------
    def cover(self) -> None:
        """The title slide: the deck's subject, its audience, its date."""
        layout = self.layouts[geometry.LAYOUT_TITLE]
        slide = self.presentation.slides.add_slide(layout)
        _set_title(slide.shapes.title, self.content.title)
        subtitle = geometry.placeholder(slide, geometry.BODY_IDX)
        lines = [line for line in (self.content.subtitle, self.context.date_line) if line]
        if not lines:
            geometry.remove_shape(subtitle)
            return
        size = _fit_single(lines, subtitle, typography.PPTX_SUBTITLE_PT)
        _fill(subtitle.text_frame, lines[:1], size)
        if len(lines) > 1:
            paragraph = subtitle.text_frame.add_paragraph()
            run = paragraph.add_run()
            run.text = lines[1]
            run.font.size = Pt(min(size, typography.PPTX_DATE_PT))
            run.font.color.rgb = RGBColor.from_string(typography.MUTED)

    def content_slide(self, spec: Slide) -> None:
        """One idea per slide; a dense one becomes "Title (k/n)" slides."""
        plans = plan_body(spec.bullets, self.body_frame)
        for index, plan in enumerate(plans, start=1):
            slide = self._new(
                geometry.LAYOUT_CONTENT,
                _continued(spec.title, index, len(plans)),
                spec.notes if index == 1 else "",
            )
            body = geometry.placeholder(slide, geometry.BODY_IDX)
            if plan.bullets:
                _fill(body.text_frame, plan.bullets, plan.size_pt)
            else:
                geometry.remove_shape(body)

    def section(self, spec: Slide) -> None:
        """A part opener on the template's own Section Header layout."""
        slide = self._new(geometry.LAYOUT_SECTION, spec.title, spec.notes)
        for paragraph in slide.shapes.title.text_frame.paragraphs:
            for run in paragraph.runs:
                # The master shouts section titles in capitals; a document does not.
                run.font._rPr.set("cap", "none")
        tagline = geometry.placeholder(slide, geometry.BODY_IDX)
        if spec.subtitle:
            size = _fit_single([spec.subtitle], tagline, typography.PPTX_SECTION_TAGLINE_PT)
            _fill(tagline.text_frame, [spec.subtitle], size, color=typography.MUTED)
        else:
            geometry.remove_shape(tagline)

    def comparison(self, spec: Slide) -> None:
        """Two sides on the Comparison layout, each split independently if needed."""
        plans = [plan_body(column.bullets, self.half_frame) for column in spec.columns]
        total = max(len(side) for side in plans)
        for index in range(total):
            slide = self._new(
                geometry.LAYOUT_COMPARISON,
                _continued(spec.title, index + 1, total),
                spec.notes if index == 0 else "",
            )
            for side, (heading_idx, body_idx) in enumerate(geometry.COMPARISON_SIDES):
                self._side(slide, spec, side, heading_idx, body_idx, plans[side], index)

    def _side(
        self,
        slide: Any,
        spec: Slide,
        side: int,
        heading_idx: int,
        body_idx: int,
        plans: list[SlidePlan],
        index: int,
    ) -> None:
        heading_shape = geometry.placeholder(slide, heading_idx)
        heading = spec.columns[side].heading
        size = _fit_single([heading], heading_shape, typography.PPTX_COMPARISON_HEADING_PT)
        _fill(heading_shape.text_frame, [heading], size)
        body = geometry.placeholder(slide, body_idx)
        plan = plans[index] if index < len(plans) else None
        if plan is not None and plan.bullets:
            _fill(body.text_frame, plan.bullets, plan.size_pt)
        else:
            geometry.remove_shape(body)

    def table(self, spec: Slide) -> None:
        """Data as native tables, chunked, header repeated on every part."""
        sheet = spec.table
        if sheet is None:
            return
        types = pptx_tables.column_types(sheet)
        chunks = pptx_tables.split_rows(sheet)
        for index, rows in enumerate(chunks, start=1):
            slide = self._new(
                geometry.LAYOUT_TITLE_ONLY,
                _continued(spec.title, index, len(chunks)),
                spec.notes if index == 1 else "",
            )
            chunk = sheet.model_copy(update={"rows": rows})
            pptx_tables.add_table(slide, chunk, types)


_KINDS: dict[str, Callable[[_Deck, Slide], None]] = {
    "content": _Deck.content_slide,
    "section": _Deck.section,
    "comparison": _Deck.comparison,
    "table": _Deck.table,
}

# Boot-time completeness (ADR-085): every slide kind the schema allows is drawn.
assert set(_KINDS) == SLIDE_KINDS, "_KINDS must cover every slide kind"


def render_pptx(content: DocumentContent, context: RenderContext) -> bytes:
    """A 16:9 deck whose every slide is the shape its content deserves."""
    if not isinstance(content, SlideContent):
        raise ValueError("pptx rendering requires SlideContent")
    deck = _Deck(content, context)
    deck.cover()
    for spec in content.slides:
        _KINDS[spec.kind](deck, spec)
    buf = io.BytesIO()
    deck.presentation.save(buf)
    return buf.getvalue()

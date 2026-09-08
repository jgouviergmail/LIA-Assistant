"""16:9 by scaling the default template; placeholders; slide numbers (ADR-274).

python-pptx's default template is 4:3, and its layout placeholders do NOT
follow a slide-size change. Scaling every master shape and every layout
placeholder that OWNS an ``xfrm`` by 4/3 gives a 13.333 × 7.5 in deck whose
placeholders sit where PowerPoint expects them — measured 2026-09-08: no
overflow, slide numbers rendered.

Writing an INHERITED position is the trap: a placeholder without its own
transform reports ``y``/``cy`` as 0, and assigning them freezes the shape at
the top of the slide with no height. Only owned transforms are touched.
"""

from __future__ import annotations

import copy
from typing import Any

import pptx
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches

from src.domains.document_generation import typography
from src.domains.document_generation.fit import BULLET_INDENT_PT, TextFrame

#: Layout indices of the default template (11 layouts, measured).
LAYOUT_TITLE = 0
LAYOUT_CONTENT = 1
LAYOUT_SECTION = 2
LAYOUT_COMPARISON = 4
LAYOUT_TITLE_ONLY = 5
#: English Metric Units per point.
EMU_PER_PT = 12700
#: Where a table sits on a "Title Only" slide, and how much room it has.
TABLE_LEFT_IN = 0.667
TABLE_TOP_IN = 1.75
TABLE_WIDTH_IN = 12.0
TABLE_AREA_HEIGHT_IN = 4.95
#: Placeholder indices of the Comparison layout: (heading, body) per side.
COMPARISON_SIDES: tuple[tuple[int, int], ...] = ((1, 2), (3, 4))
#: The body placeholder index on Content, Section Header and the two-content layouts.
BODY_IDX = 1


def _scale_owned(shapes: Any, scale: float) -> None:
    """Scale the shapes that carry their own transform; leave inherited ones alone."""
    for shape in shapes:
        shape_properties = shape._element.spPr
        if shape_properties is None or shape_properties.xfrm is None:
            continue
        shape.left = int(shape.left * scale)
        shape.width = int(shape.width * scale)


def base_presentation() -> Any:
    """The default template turned 16:9 LANDSCAPE, placeholders in place."""
    presentation = pptx.Presentation()
    width = Inches(typography.PPTX_SLIDE_WIDTH_IN)
    scale = width / presentation.slide_width
    presentation.slide_width = width
    presentation.slide_height = Inches(typography.PPTX_SLIDE_HEIGHT_IN)
    _scale_owned(presentation.slide_master.shapes, scale)
    for layout in presentation.slide_layouts:
        _scale_owned(layout.placeholders, scale)
    return presentation


def placeholder(container: Any, idx: int) -> Any:
    """The placeholder with this layout index, on a slide or on a layout.

    Args:
        container: A slide or a slide layout (both expose ``placeholders``).
        idx: The placeholder index.

    Returns:
        The placeholder shape.

    Raises:
        KeyError: When the container has no such placeholder.
    """
    for candidate in container.placeholders:
        if candidate.placeholder_format.idx == idx:
            return candidate
    raise KeyError(idx)


def add_slide_number(slide: Any, layout: Any) -> None:
    """Clone the layout's slide-number placeholder; PowerPoint renders the field."""
    for candidate in layout.placeholders:
        if candidate.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            slide.shapes._spTree.append(copy.deepcopy(candidate._element))
            return


def remove_shape(shape: Any) -> None:
    """Drop an unused placeholder, so no "Click to add text" prompt survives."""
    element = shape._element
    element.getparent().remove(element)


def frame_of(shape: Any, indent_pt: float = BULLET_INDENT_PT) -> TextFrame:
    """The shape's box as a text frame in points."""
    return TextFrame(shape.width / EMU_PER_PT, shape.height / EMU_PER_PT, indent_pt=indent_pt)

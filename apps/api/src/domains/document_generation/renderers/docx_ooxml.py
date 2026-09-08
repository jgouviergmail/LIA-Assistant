"""Raw OOXML python-docx has no API for (ADR-274): fields, numbering, borders.

Word computes some things itself — the page count, the table of contents, the
heading numbers — and a document that WRITES those as text is a document whose
numbers are wrong the moment a paragraph is added. Each helper here does ONE
thing to ONE element, so the renderer stays readable and every mechanism is
exercised by the docx tests through python-docx's own reader.
"""

from __future__ import annotations

from datetime import UTC, datetime
from typing import Any

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

#: Levels the multilevel definition numbers (1 / 1.1 / 1.1.1).
NUMBERED_LEVELS = 3
#: Border width of the callout's left rule, in eighths of a point.
_CALLOUT_BORDER_SIZE = "18"


def naive_utc(moment: datetime) -> datetime:
    """Core properties are written as naive UTC by python-docx and openpyxl alike."""
    return moment.astimezone(UTC).replace(tzinfo=None)


def _element(tag: str, **attributes: str) -> Any:
    """One OOXML element with ``w:``-namespaced attributes."""
    element = OxmlElement(tag)
    for name, value in attributes.items():
        element.set(qn(f"w:{name}"), value)
    return element


def _instruction(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    run._r.append(text)


def open_field(paragraph: Any, instruction: str) -> None:
    """The begin/instruction/separate triple of a field whose result spans paragraphs."""
    paragraph.add_run()._r.append(_element("w:fldChar", fldCharType="begin"))
    _instruction(paragraph, instruction)
    paragraph.add_run()._r.append(_element("w:fldChar", fldCharType="separate"))


def close_field(paragraph: Any) -> None:
    """The end marker of a field opened with :func:`open_field`."""
    paragraph.add_run()._r.append(_element("w:fldChar", fldCharType="end"))


def add_field(paragraph: Any, instruction: str, cached: str = "") -> None:
    """A complete field: begin / instruction / separate / cached result / end.

    Args:
        paragraph: Where the field goes.
        instruction: The field code, e.g. ``PAGE`` or ``NUMPAGES``.
        cached: What a reader sees before Word recomputes the field.
    """
    open_field(paragraph, instruction)
    paragraph.add_run(cached)
    close_field(paragraph)


def mark_header_row(row: Any) -> None:
    """Repeat this row at the top of every page the table continues on."""
    row._tr.get_or_add_trPr().append(_element("w:tblHeader", val="true"))


def set_table_look(table: Any, *, first_column: bool) -> None:
    """Which conditional formats of the style apply (header row on, banding on)."""
    table_properties = table._tbl.tblPr
    for stale in table_properties.findall(qn("w:tblLook")):
        table_properties.remove(stale)
    table_properties.append(
        _element(
            "w:tblLook",
            firstRow="1",
            lastRow="0",
            firstColumn="1" if first_column else "0",
            lastColumn="0",
            noHBand="0",
            noVBand="1",
        )
    )


def _numbering_root(document: Any) -> Any:
    return document.part.numbering_part.element


def _next_id(root: Any, tag: str, attribute: str) -> int:
    return max((int(node.get(qn(attribute))) for node in root.findall(qn(tag))), default=-1) + 1


def add_heading_numbering(document: Any) -> None:
    """One multilevel list bound to Heading 1-3, so Word numbers them itself."""
    root = _numbering_root(document)
    abstract_id = _next_id(root, "w:abstractNum", "w:abstractNumId")
    abstract = _element("w:abstractNum", abstractNumId=str(abstract_id))
    abstract.append(_element("w:multiLevelType", val="multilevel"))
    for level, pattern in enumerate(("%1", "%1.%2", "%1.%2.%3")):
        definition = _element("w:lvl", ilvl=str(level))
        definition.append(_element("w:start", val="1"))
        definition.append(_element("w:numFmt", val="decimal"))
        definition.append(_element("w:pStyle", val=f"Heading{level + 1}"))
        definition.append(_element("w:lvlText", val=pattern))
        definition.append(_element("w:lvlJc", val="left"))
        paragraph_properties = OxmlElement("w:pPr")
        paragraph_properties.append(_element("w:ind", left="0", hanging="0"))
        definition.append(paragraph_properties)
        abstract.append(definition)
    first_num = root.find(qn("w:num"))
    if first_num is not None:
        # abstractNum elements must precede num elements, or Word repairs the file.
        first_num.addprevious(abstract)
    else:
        root.append(abstract)
    num_id = max(_next_id(root, "w:num", "w:numId"), 1)
    num = _element("w:num", numId=str(num_id))
    num.append(_element("w:abstractNumId", val=str(abstract_id)))
    root.append(num)
    for level in range(NUMBERED_LEVELS):
        properties = document.styles[f"Heading {level + 1}"].element.get_or_add_pPr()
        numbering = OxmlElement("w:numPr")
        numbering.append(_element("w:ilvl", val=str(level)))
        numbering.append(_element("w:numId", val=str(num_id)))
        properties.append(numbering)


def detach_toc_heading(document: Any) -> None:
    """``TOC Heading`` is based on Heading 1: take it out of numbering and outline.

    Measured: without this, "Contents" is counted as chapter 1 and appears in
    the table of contents it introduces.
    """
    properties = document.styles["TOC Heading"].element.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    numbering.append(_element("w:ilvl", val="0"))
    numbering.append(_element("w:numId", val="0"))
    properties.append(numbering)
    properties.append(_element("w:outlineLvl", val="9"))


def fresh_list_num(document: Any, style_name: str) -> int:
    """A new numbering instance on the style's list, restarting at 1.

    Args:
        document: The document being written.
        style_name: A numbered style, e.g. ``List Number``.

    Returns:
        The new ``numId`` to bind the list's paragraphs to.
    """
    root = _numbering_root(document)
    style_properties = document.styles[style_name].element.pPr
    style_num_id = style_properties.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
    abstract_id = next(
        num.find(qn("w:abstractNumId")).get(qn("w:val"))
        for num in root.findall(qn("w:num"))
        if num.get(qn("w:numId")) == style_num_id
    )
    num_id = _next_id(root, "w:num", "w:numId")
    num = _element("w:num", numId=str(num_id))
    num.append(_element("w:abstractNumId", val=abstract_id))
    override = _element("w:lvlOverride", ilvl="0")
    override.append(_element("w:startOverride", val="1"))
    num.append(override)
    root.append(num)
    return num_id


def set_paragraph_num(paragraph: Any, num_id: int) -> None:
    """Bind a list paragraph to a numbering instance."""
    numbering = OxmlElement("w:numPr")
    numbering.append(_element("w:ilvl", val="0"))
    numbering.append(_element("w:numId", val=str(num_id)))
    paragraph._p.get_or_add_pPr().append(numbering)


def set_east_asian_font(style: Any, name: str) -> None:
    """Declare the East-Asian face so CJK never falls back to a Latin one."""
    run_properties = style.element.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.append(fonts)
    fonts.set(qn("w:eastAsia"), name)


def shade_paragraph_style(style: Any, fill_hex: str) -> None:
    """Background fill of every paragraph of the style."""
    style.element.get_or_add_pPr().append(
        _element("w:shd", val="clear", color="auto", fill=fill_hex)
    )


def left_border_style(style: Any, color_hex: str) -> None:
    """A left rule on every paragraph of the style."""
    borders = OxmlElement("w:pBdr")
    borders.append(
        _element("w:left", val="single", sz=_CALLOUT_BORDER_SIZE, space="8", color=color_hex)
    )
    style.element.get_or_add_pPr().append(borders)


def recolor_title_rule(style: Any, color_hex: str) -> None:
    """The Title style's theme-blue bottom rule becomes a neutral grey one."""
    properties = style.element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    bottom = borders.find(qn("w:bottom")) if borders is not None else None
    if bottom is None:
        return
    bottom.set(qn("w:color"), color_hex)
    for themed in ("w:themeColor", "w:themeShade", "w:themeTint"):
        bottom.attrib.pop(qn(themed), None)

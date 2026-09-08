"""xlsx: typed columns, a named Table, a frozen header, a filter (ADR-226, ADR-274).

A spreadsheet whose numbers are strings is a picture of data, not data: it
cannot be summed, sorted or filtered. Columns are therefore TYPED (by unanimity
— see ``tables.py``) and carry their display format, the header row is frozen
and the rows sit in a named Table so Excel offers its filter.

Two rules Excel enforces and we must satisfy before writing: a Table's headers
are unique and non-empty, and its name is an identifier — both measured
2026-09-08, both handled upstream in ``tables.py``.
"""

from __future__ import annotations

import io
import re
from datetime import UTC

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from src.domains.document_generation import typography
from src.domains.document_generation.context import RenderContext
from src.domains.document_generation.fit import display_columns
from src.domains.document_generation.sanitize import neutralize_formula
from src.domains.document_generation.schemas import DocumentContent, TableSheet, TabularContent
from src.domains.document_generation.tables import (
    ColumnKind,
    ColumnType,
    excel_table_name,
    infer_column_types,
    typed_value,
)

_XLSX_TITLE_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")  # openpyxl rejects these
_XLSX_TITLE_MAX = 31  # Excel's hard sheet-title limit
#: Padding added to the longest cell when sizing a column.
_WIDTH_PADDING = 2


def _xlsx_sheet_title(name: str, index: int, used: set[str]) -> str:
    """Sanitize an LLM-suggested worksheet title for openpyxl.

    Args:
        name: Suggested sheet name.
        index: Zero-based sheet position (for the fallback name).
        used: Titles already taken in this workbook (mutated in place).

    Returns:
        A unique, openpyxl-legal worksheet title.
    """
    cleaned = _XLSX_TITLE_FORBIDDEN.sub("_", name).strip()[:_XLSX_TITLE_MAX]
    cleaned = cleaned or f"Sheet{index + 1}"
    candidate = cleaned
    suffix = 2
    while candidate in used:
        tail = f" {suffix}"
        candidate = f"{cleaned[: _XLSX_TITLE_MAX - len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate


def _alignment_for(column: ColumnType, width: int) -> Alignment | None:
    """How a column's cells sit: numbers right, dates centred, long text wrapped."""
    if column.numeric:
        return Alignment(horizontal="right")
    if column.kind in (ColumnKind.DATE, ColumnKind.DATETIME):
        return Alignment(horizontal="center")
    if width >= typography.XLSX_COLUMN_WIDTH_MAX:
        return Alignment(wrap_text=True, vertical="top")
    return None


def _column_width(header: str, rows: list[list[str]], index: int) -> int:
    """Width in Excel's own unit: the average Latin character, not the code point."""
    longest = max(
        [display_columns(header), *(display_columns(row[index]) for row in rows)],
        default=display_columns(header),
    )
    return min(
        max(int(longest) + _WIDTH_PADDING, typography.XLSX_COLUMN_WIDTH_MIN),
        typography.XLSX_COLUMN_WIDTH_MAX,
    )


def _write_header(worksheet: Worksheet, headers: list[str]) -> None:
    worksheet.append([neutralize_formula(header) for header in headers])
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _add_table(worksheet: Worksheet, sheet: TableSheet, table_index: int) -> None:
    """The named Table that gives Excel its filter and its banding."""
    reference = f"A1:{get_column_letter(len(sheet.headers))}{len(sheet.rows) + 1}"
    table = Table(displayName=excel_table_name(table_index), ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name=typography.XLSX_TABLE_STYLE,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _fill_sheet(worksheet: Worksheet, sheet: TableSheet, table_index: int) -> None:
    """One worksheet: typed cells, sized columns, frozen header, Table."""
    types = infer_column_types(sheet.headers, sheet.rows)
    _write_header(worksheet, sheet.headers)
    for row in sheet.rows:
        worksheet.append(
            [typed_value(value, column) for value, column in zip(row, types, strict=True)]
        )
    for index, (header, column) in enumerate(zip(sheet.headers, types, strict=True), start=1):
        letter = get_column_letter(index)
        width = _column_width(header, sheet.rows, index - 1)
        worksheet.column_dimensions[letter].width = width
        alignment = _alignment_for(column, width)
        for row_index in range(2, len(sheet.rows) + 2):
            cell = worksheet.cell(row=row_index, column=index)
            cell.number_format = column.number_format
            if alignment is not None:
                cell.alignment = alignment
    worksheet.freeze_panes = "A2"
    if sheet.rows:
        # A Table over a header-only range is what Excel refuses to open.
        _add_table(worksheet, sheet, table_index)


def render_xlsx(content: DocumentContent, context: RenderContext) -> bytes:
    """A workbook whose columns are data, not a picture of data."""
    if not isinstance(content, TabularContent):
        raise ValueError("xlsx rendering requires TabularContent")
    workbook = Workbook()
    workbook.properties.title = content.title
    if context.generated_at is not None:
        workbook.properties.created = context.generated_at.astimezone(UTC).replace(tzinfo=None)
    default_sheet = workbook.active
    used_titles: set[str] = set()
    for index, sheet in enumerate(content.sheets):
        worksheet = default_sheet if index == 0 else workbook.create_sheet()
        worksheet.title = _xlsx_sheet_title(sheet.name, index, used_titles)
        _fill_sheet(worksheet, sheet, index + 1)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()

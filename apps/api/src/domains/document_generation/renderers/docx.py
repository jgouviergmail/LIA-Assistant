"""docx: named styles, fields, numbering, table of contents (ADR-226, ADR-274).

Everything is a NAMED style (``docx_styles``), and everything Word computes
itself — page numbers, the table of contents, heading numbers — is a field or a
numbering definition rather than text pretending to be one.

The long-document apparatus follows ONE predicate (``document_is_numbered``):
table of contents, heading numbering and a page break before each part switch
on together, so a two-page note never grows an apparatus it does not need.

Word's table of contents is PRE-RENDERED without page numbers (the ``\\n``
switch): the field is empty until Word recomputes it, so a document opened and
printed straight away would otherwise show a blank contents page. F9 regenerates
it identically.
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from typing import Any

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Mm, RGBColor

from src.domains.document_generation import typography
from src.domains.document_generation.context import RenderContext
from src.domains.document_generation.inline import parse_inline
from src.domains.document_generation.normalize import HeadingNumberer, document_is_numbered
from src.domains.document_generation.renderers import docx_ooxml as ooxml
from src.domains.document_generation.renderers.docx_styles import CALLOUT_STYLE, define_styles
from src.domains.document_generation.schemas import (
    BLOCK_KINDS,
    DocumentContent,
    SectionBlock,
    SectionedContent,
)
from src.domains.document_generation.tables import infer_column_types

#: Splits the localized "Page {page} / {total}" around its two fields.
_PAGE_OF_TOKENS = re.compile(r"(\{page\}|\{total\})")
_TOC_FIELD = 'TOC \\o "1-{levels}" \\h \\z \\u \\n'
_LETTER_SIZE_IN = (8.5, 11)
_A4_SIZE_MM = (210, 297)


class _Writer:
    """Per-document state: the python-docx document, its numbering, its tables."""

    def __init__(self, content: SectionedContent, context: RenderContext) -> None:
        self.content = content
        self.context = context
        self.document = docx.Document()
        self.numbered = document_is_numbered(content, context)
        # No numberer here on purpose: WORD numbers the headings from the
        # numbering definition, and the contents entries count with their own
        # (a shared counter would number the contents and the body in one run).
        self.tables = 0
        self.first_part_seen = False

    # -- setup ------------------------------------------------------------
    def setup(self) -> None:
        """Page, styles and document properties."""
        section = self.document.sections[0]
        if self.context.page_size == "letter":
            section.page_width, section.page_height = (Inches(n) for n in _LETTER_SIZE_IN)
        else:
            section.page_width, section.page_height = (Mm(n) for n in _A4_SIZE_MM)
        margin = Cm(typography.docx_margin_cm(self.context.page_size))
        section.left_margin = section.right_margin = margin
        section.top_margin = section.bottom_margin = margin
        define_styles(self.document, numbered=self.numbered)
        properties = self.document.core_properties
        properties.title = self.content.title
        properties.subject = self.content.subtitle
        properties.language = self.context.language
        if self.context.generated_at is not None:
            properties.created = ooxml.naive_utc(self.context.generated_at)

    # -- front matter -----------------------------------------------------
    def front_matter(self) -> None:
        """Title block, running head, page footer, and the contents when long."""
        self.document.add_paragraph(self.content.title, style="Title")
        if self.content.subtitle:
            self.document.add_paragraph(self.content.subtitle, style="Subtitle")
        if self.context.date_line:
            date = self.document.add_paragraph()
            date.add_run(self.context.date_line).font.color.rgb = RGBColor.from_string(
                typography.MUTED
            )
        section = self.document.sections[0]
        header = section.header.paragraphs[0]
        header.text = self.content.title
        header.style = self.document.styles["Header"]
        self._footer(section)
        if self.numbered:
            self._table_of_contents()

    def _footer(self, section: Any) -> None:
        """« Page 2 / 7 » from PAGE and NUMPAGES, in the reader's word order."""
        footer = section.footer.paragraphs[0]
        footer.style = self.document.styles["Footer"]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        template = self.context.label("documents.page_of", page="{page}", total="{total}")
        for piece in _PAGE_OF_TOKENS.split(template):
            if piece == "{page}":
                ooxml.add_field(footer, "PAGE", "1")
            elif piece == "{total}":
                ooxml.add_field(footer, "NUMPAGES", "1")
            elif piece:
                footer.add_run(piece)

    def _toc_entries(self) -> list[tuple[int, str]]:
        numberer = HeadingNumberer()
        return [
            (block.level, f"{numberer.number(block.level)}\t{block.text}")
            for block in self.content.blocks
            if block.kind == "heading" and block.level <= ooxml.NUMBERED_LEVELS
        ]

    def _table_of_contents(self) -> None:
        self.document.add_paragraph(
            self.context.label("documents.toc_heading"), style="TOC Heading"
        )
        paragraph = None
        for index, (level, text) in enumerate(self._toc_entries()):
            paragraph = self.document.add_paragraph(style=f"TOC {level}")
            if index == 0:
                ooxml.open_field(paragraph, _TOC_FIELD.format(levels=ooxml.NUMBERED_LEVELS))
            paragraph.add_run(text)
        if paragraph is not None:
            ooxml.close_field(paragraph)
        self.document.add_page_break()

    # -- blocks -----------------------------------------------------------
    def runs(self, paragraph: Any, text: str) -> None:
        """The model's inline emphasis as Word runs."""
        for span in parse_inline(text):
            run = paragraph.add_run(span.text)
            run.bold = span.bold or None
            run.italic = span.italic or None
            if span.code:
                run.font.name = typography.FONT_MONO

    def heading(self, block: SectionBlock) -> None:
        paragraph = self.document.add_heading(block.text, level=block.level)
        if self.numbered and block.level == 1:
            paragraph.paragraph_format.page_break_before = self.first_part_seen
            self.first_part_seen = True

    def paragraph(self, block: SectionBlock) -> None:
        self.runs(self.document.add_paragraph(), block.text)

    def bullets(self, block: SectionBlock) -> None:
        for item in block.items:
            self.runs(self.document.add_paragraph(style="List Bullet"), item)

    def numbered_list(self, block: SectionBlock) -> None:
        """A fresh numbering instance per list, so every sequence restarts at 1."""
        num_id = ooxml.fresh_list_num(self.document, "List Number")
        for item in block.items:
            paragraph = self.document.add_paragraph(style="List Number")
            ooxml.set_paragraph_num(paragraph, num_id)
            self.runs(paragraph, item)

    def quote(self, block: SectionBlock) -> None:
        self.runs(self.document.add_paragraph(style="Quote"), block.text)

    def callout(self, block: SectionBlock) -> None:
        self.runs(self.document.add_paragraph(style=CALLOUT_STYLE), block.text)

    def table(self, block: SectionBlock) -> None:
        """A captioned table whose header repeats on every page it spans."""
        sheet = block.table
        if sheet is None:
            return
        self.tables += 1
        if block.caption:
            label = self.context.label("documents.table_label", n=self.tables)
            self.document.add_paragraph(f"{label} — {block.caption}", style="Caption")
        types = infer_column_types(sheet.headers, sheet.rows)
        table = self.document.add_table(rows=1, cols=len(sheet.headers))
        table.style = typography.DOCX_TABLE_STYLE
        table.autofit = True
        ooxml.set_table_look(table, first_column=False)
        for index, header in enumerate(sheet.headers):
            table.cell(0, index).text = header
        ooxml.mark_header_row(table.rows[0])
        for row in sheet.rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = value
                if types[index].numeric:
                    cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.document.add_paragraph()


_BLOCKS: dict[str, Callable[[_Writer, SectionBlock], None]] = {
    "heading": _Writer.heading,
    "paragraph": _Writer.paragraph,
    "bullets": _Writer.bullets,
    "numbered": _Writer.numbered_list,
    "quote": _Writer.quote,
    "callout": _Writer.callout,
    "table": _Writer.table,
}

# Boot-time completeness (ADR-085): every block kind the schema allows is drawn.
assert set(_BLOCKS) == BLOCK_KINDS, "_BLOCKS must cover every block kind"


def render_docx(content: DocumentContent, context: RenderContext) -> bytes:
    """A Word document a reader can restyle, navigate and print."""
    if not isinstance(content, SectionedContent):
        raise ValueError("docx rendering requires SectionedContent")
    writer = _Writer(content, context)
    writer.setup()
    writer.front_matter()
    for block in content.blocks:
        _BLOCKS[block.kind](writer, block)
    buf = io.BytesIO()
    writer.document.save(buf)
    return buf.getvalue()
